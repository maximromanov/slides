import sys, json
from pptx import Presentation
src, out = sys.argv[1], sys.argv[2]
prs = Presentation(src); tables = {}
for i, s in enumerate(prs.slides, 1):
    for sh in s.shapes:
        if sh.has_table:
            rows = [[c.text.strip() for c in r.cells] for r in sh.table.rows]
            tables.setdefault(i, []).append(rows)
json.dump(tables, open(out, 'w'), ensure_ascii=False, indent=1)
print(sys.argv[1][-30:], 'tables on slides:', sorted(tables))
