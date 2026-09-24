import sys, json, hashlib, os
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from PIL import Image
src, out = sys.argv[1], sys.argv[2]
os.makedirs(out+'/img', exist_ok=True)
prs = Presentation(src); W, H = prs.slide_width, prs.slide_height
pct = lambda v, t: round((v or 0)/t*100, 1)
slides = []; seen = {}
for i, s in enumerate(prs.slides, 1):
    e = {'n': i, 'shapes': [], 'notes': s.notes_slide.notes_text_frame.text.strip() if s.has_notes_slide else ''}
    for sh in s.shapes:
        d = {'x': pct(sh.left, W), 'y': pct(sh.top, H), 'w': pct(sh.width, W), 'h': pct(sh.height, H)}
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            img = sh.image; dg = hashlib.md5(img.blob).hexdigest()[:10]; fn = f'img/{dg}.{img.ext}'
            if dg not in seen:
                open(f'{out}/{fn}', 'wb').write(img.blob); seen[dg] = fn
            d['image'] = fn
            try: d['crop'] = [round(sh.crop_left,2), round(sh.crop_top,2), round(sh.crop_right,2), round(sh.crop_bottom,2)]
            except Exception: pass
        if sh.has_text_frame and sh.text_frame.text.strip():
            d['text'] = [{'t': p.text, 'lvl': p.level, 'size': next((r.font.size.pt for r in p.runs if r.font.size), None), 'b': any(r.font.bold for r in p.runs)} for p in sh.text_frame.paragraphs if p.text.strip()]
        if sh.shape_type == MSO_SHAPE_TYPE.GROUP: d['group'] = True
        e['shapes'].append(d)
    slides.append(e)
json.dump(slides, open(out+'/slides.json','w'), ensure_ascii=False, indent=1)
print(os.path.basename(src), len(slides), 'slides,', len(seen), 'images')
